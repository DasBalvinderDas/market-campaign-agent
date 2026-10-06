#!/usr/bin/env python3
"""Builds docs/Campaign_Provisioner_Agent_Catalog.xlsx: root agent and sub-agents, tools, BigQuery and Application Integration.

  python docs/build_agent_catalog.py        (needs: pip install openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "Campaign_Provisioner_Agent_Catalog.xlsx"
APP, OWNER = "Campaign Provisioner (Google Next 2027 Marketing Logistics Agent)", "Dasbalvinder Das"

AGENTS = [
    ("Root Agent", "Campaign Provisioner (Orchestrator)",
     "The root agent is the front-line coordinator. It greets users, collects the campaign and the items, registers the request, and "
     "calls the three specialist agents in a fixed order (inventory, procurement, budget). It holds the human-in-the-loop approval gate: "
     "only the root can call approve_budget. Amounts up to $5,000 are approved by policy; higher amounts start an Application Integration "
     "approval and the reply opens with HUMAN APPROVAL REQUIRED (nothing is committed or ordered until the approver says yes). When the user "
     "asks for the approval status it reads the decision, then commits the budget and creates the purchase orders (Approved) or releases the "
     "stock (Rejected). It writes every reply the user sees.",
     "BigQuery (requests, budget ledger, approvals, audit log); Application Integration (request_approval, create_purchase_order); "
     "Vertex AI Agent Engine; Gemini Enterprise chat"),
    ("Sub Agent", "Inventory Agent",
     "The inventory agent matches each requested item (free text or SKU) to the catalog, checks free stock and reserves what is available, "
     "for all items in one call. It records every line (requested, reserved, shortfall) so the later steps know what has to be bought, and "
     "it releases the reserved stock when a request is rejected or declined. It never buys anything and never discusses budget.",
     "BigQuery: inventory_items / v_inventory_available, inventory_reservations, request_lines, audit_log"),
    ("Sub Agent", "Procurement Agent",
     "The procurement agent quotes every shortfall line at once from the cheapest active vendor (price and lead time) and reports the line "
     "totals and the grand total. After the budget is approved it creates one purchase order per line through the Application Integration "
     "create_purchase_order trigger. A code guard wraps that call: it recomputes the amount from the vendor price list and blocks the order "
     "unless an approved budget covers it. It never approves budget.",
     "BigQuery: request_lines, vendor_catalog, vendors, v_request_headroom, purchase_orders, audit_log; "
     "Application Integration: create_purchase_order"),
    ("Sub Agent", "Budget Agent",
     "The budget agent checks the campaign's remaining budget and finds the approval tier for the total (AUTO up to $5,000, MANAGER up to "
     "$50,000, EXECUTIVE above) and the approver role, in one call, and reports back to the root. It never approves spend and never places "
     "orders. (With the chat approval channel it also emails the approver through the notify_approver trigger.)",
     "BigQuery: v_campaign_budget, approval_policy, approvers; Application Integration: notify_approver (chat channel only)"),
]

TOOLS = [
    # agent, tool, kind, what it does, BQ reads, BQ writes, Application Integration, guardrail
    ("Root Agent", "register_campaign_request", "Python tool", "Validates the campaign and creates the request (request id REQ-xxxx).", "campaigns / v_campaign_budget", "campaign_requests, audit_log", "-", "Unknown campaign is refused"),
    ("Root Agent", "inventory_agent / procurement_agent / budget_agent", "Agent as tool (ADK AgentTool)", "The root calls each specialist like a tool, passing the request id, items and amounts; the specialist runs with its own prompt and tools and returns its result.", "(via the specialist)", "(via the specialist)", "(via the specialist)", "Fixed order in the root instruction; specialists reply only to the root"),
    ("Root Agent", "approve_budget", "Python tool (human gate)", "Decides the budget. AUTO tier: commits the budget at once. Human tier: builds the purchase plan, stores a PENDING approval and starts the Application Integration approval; nothing is committed. Over budget: rejected.", "v_campaign_budget, approval_policy, request_lines, vendor_catalog, vendors, approvers, campaign_requests", "budget_ledger (COMMIT, auto tier), approval_requests (PENDING), audit_log", "request_approval (started over REST)", "Only the root has this tool; the amount comes from BigQuery, not from the model; overspend is rejected"),
    ("Root Agent", "get_approval_status", "Python tool", "Reads the approval. If the approver decided, records it (first decision wins) and carries it out: Approved = commit budget + create the purchase orders; Rejected = release the reserved stock.", "approval_requests, purchase_orders, vendor_catalog, v_request_headroom", "approval_requests (status), budget_ledger (COMMIT), purchase_orders, inventory_reservations (RELEASED), audit_log", "Reads the execution and its approval record (LIFTED = approved, REJECTED = rejected); create_purchase_order per line", "Atomic update: asking twice never orders twice"),
    ("Root Agent", "get_campaign_overview", "Python tool", "Budget position of every campaign (total, spent, committed, remaining).", "v_campaign_budget", "-", "-", "Read only"),
    ("Root Agent", "get_audit_trail", "Python tool", "Shows the ordered audit events (all or for one request).", "audit_log", "-", "-", "Read only"),
    ("Inventory Agent", "reserve_items", "Python tool (batch)", "Matches every item to a catalog SKU, checks free stock, reserves the free units and records each line. Unknown or ambiguous items come back as a question for the user.", "v_inventory_available", "inventory_reservations, request_lines, audit_log", "-", "Never invents a SKU; reserves at most the free stock"),
    ("Inventory Agent", "find_sku / check_inventory / reserve_inventory", "Python tools (single item)", "One-item versions of the same steps, for a follow-up.", "v_inventory_available", "inventory_reservations, request_lines, audit_log (reserve)", "-", "Unknown SKU returns the valid SKUs"),
    ("Inventory Agent", "release_inventory", "Python tool", "Releases all stock reserved for a request (rejected or declined budget).", "inventory_reservations", "inventory_reservations (status RELEASED), audit_log", "-", "-"),
    ("Procurement Agent", "quote_shortfalls", "Python tool (batch)", "Quotes every shortfall line from the cheapest active vendor with price, lead time and line totals, plus the grand total.", "request_lines, vendor_catalog, vendors", "-", "-", "Read only; unfulfillable SKUs are listed"),
    ("Procurement Agent", "get_vendor_quotes", "Python tool (single item)", "Vendor quotes for one SKU or item description, cheapest first.", "vendor_catalog, vendors", "-", "-", "Resolves a description to a SKU through the catalog"),
    ("Procurement Agent", "create_purchase_order", "Application Integration tool (ADK ApplicationIntegrationToolset)", "Creates a purchase order: the integration sets the PO number (PO-<request>-<sku>) and returns the execution id.", "Guard reads vendor_catalog and v_request_headroom", "purchase_orders, audit_log (after-tool callback)", "create_purchase_order", "Before-tool guard recomputes the amount and returns BLOCKED unless an approved budget covers it"),
    ("Budget Agent", "assess_budget", "Python tool (batch)", "Remaining budget, whether the total fits, the approval tier, whether a human is needed and the approver role, in one call.", "v_campaign_budget, approval_policy", "-", "-", "Read only"),
    ("Budget Agent", "check_budget / get_approval_policy", "Python tools (single)", "The two halves of assess_budget.", "v_campaign_budget / approval_policy", "-", "-", "Read only"),
    ("Budget Agent", "notify_approver", "Application Integration tool (chat channel only)", "Emails one approver address; the guard fills the recipient, subject and body from BigQuery, one call per address.", "approvers", "audit_log", "notify_approver", "The model never supplies an email address"),
]

BQ = [
    ("inventory_items", "Table", "Catalog and stock on hand (sku, name, category, on_hand).", "Inventory agent", "Setup script (seed)"),
    ("inventory_reservations", "Table", "Stock reserved per request (ACTIVE or RELEASED).", "Inventory agent (view)", "Inventory agent, root (release)"),
    ("v_inventory_available", "View", "Free stock = on hand minus active reservations.", "Inventory agent", "-"),
    ("vendors", "Table", "Vendors and whether they are active.", "Procurement agent, guard", "Setup script (seed)"),
    ("vendor_catalog", "Table", "Vendor price and lead time per SKU.", "Procurement agent, root, guard", "Setup script (seed)"),
    ("campaigns", "Table", "Campaigns (NEXT27-MAIN, NEXT27-PARTNER, NEXT27-DEVLOUNGE) and their total budget.", "Root, budget agent", "Setup script (seed)"),
    ("budget_ledger", "Table", "Budget ledger: SPEND (already spent, BASELINE), COMMIT (approved by the agent), RELEASE.", "v_campaign_budget", "Root (approve_budget, get_approval_status)"),
    ("v_campaign_budget", "View", "Total, spent, committed and remaining per campaign.", "Root, budget agent", "-"),
    ("v_request_headroom", "View", "Approved amount vs. amount already ordered per request (used by the PO guard).", "Guard, root", "-"),
    ("approval_policy", "Table", "Approval tiers: amount range, human needed, approver role.", "Budget agent, root", "Setup script (seed); editable by finance"),
    ("approvers", "Table", "Approver email addresses per role (configurable).", "Root, guard", "Setup script (--approver-email)"),
    ("request_lines", "Table", "Per request and SKU: requested, reserved, shortfall (the purchase plan source).", "Procurement agent, root", "Inventory agent"),
    ("approval_requests", "Table", "Human approvals: approver, amount, plan, status (PENDING / APPROVED / REJECTED / FAILED), workflow execution id.", "Root (get_approval_status)", "Root (approve_budget, get_approval_status)"),
    ("campaign_requests", "Table", "Every request registered by the root agent.", "Root", "Root (register_campaign_request)"),
    ("purchase_orders", "Table", "Purchase orders with the Application Integration execution id.", "Root, status checks", "Procurement guard (after-tool), root (after approval)"),
    ("audit_log", "Table", "Every governed action (who, what, when, details).", "Root (get_audit_trail)", "All agents' tools and guards"),
]

AI = [
    ("create_purchase_order", "API trigger", "Create a purchase order.", "request_id, campaign_id, sku, vendor_id, quantity, total_amount", "po_number (PO-<request_id>-<sku>), execution_id", "Data Mapping task: builds the PO number and execution id.", "Procurement agent (ADK Application Integration tool, behind the guard); root after an approval (REST)", "After the budget is approved"),
    ("request_approval", "API trigger (native approval)", "Human approval: the run pauses until an approver decides.", "request_id, campaign_id, approver_role, amount, approval_message", "decision (APPROVED / REJECTED)", "Approval task (suspension): emails the approver(s) an Approve / Reject request, reminder after 1 day, expires after 3 days; branch isApproved = true -> Data Mapping decision = APPROVED; branch false -> decision = REJECTED. The click is recorded as an approval record (LIFTED = approved, REJECTED = rejected).", "Root (approve_budget starts it over REST; get_approval_status reads the result)", "Amount above $5,000"),
    ("notify_approver", "API trigger", "Send one plain notification email.", "request_id, campaign_id, approver_role, approver_email, email_subject, email_body, summary, amount", "status (NOTIFIED)", "Send Email task (one address per call), then Data Mapping sets status.", "Budget agent (chat approval channel only); also used by the setup test", "Only when the approval channel is chat"),
]

FLOW = [
    (1, "User", "Sends the request in Gemini Enterprise (or the playground).", "-", "-", "NEXT27-MAIN needs 1 booth LED video wall and 8 event banners."),
    (2, "Root Agent", "register_campaign_request", "Reads campaigns; writes campaign_requests, audit_log", "-", "Request id REQ-xxxx"),
    (3, "Inventory Agent", "reserve_items (one call)", "Reads v_inventory_available; writes inventory_reservations, request_lines, audit_log", "-", "2 banners reserved, 6 short; LED wall 1 short"),
    (4, "Procurement Agent", "quote_shortfalls (one call)", "Reads request_lines, vendor_catalog, vendors", "-", "LED wall ExpoVision $18,000, 6 banners PrintCo $510: total $18,510"),
    (5, "Budget Agent", "assess_budget (one call)", "Reads v_campaign_budget, approval_policy", "-", "Fits the budget; tier MANAGER; approver Marketing Director"),
    (6, "Root Agent", "approve_budget", "Reads plan, approvers; writes approval_requests (PENDING), audit_log", "Starts request_approval (Approval task pauses; approver emailed)", "Reply opens with HUMAN APPROVAL REQUIRED; nothing committed, no PO"),
    (7, "Approver", "Opens the Application Integration email and clicks Approve (or Reject)", "-", "Approval record becomes LIFTED (approve) or REJECTED", "Decision recorded in Application Integration"),
    (8, "User", "Asks: What is the approval status of REQ-xxxx?", "-", "-", "-"),
    (9, "Root Agent", "get_approval_status", "Reads approval_requests; Approved: writes budget_ledger COMMIT, approval_requests APPROVED, audit_log", "Reads the execution and approval record", "-"),
    (10, "Procurement Agent / root", "Create the purchase orders (Approved)", "Guard reads vendor_catalog, v_request_headroom; writes purchase_orders, audit_log", "create_purchase_order per line", "PO numbers for the LED wall and the banners; remaining budget $231,490"),
    (11, "Root Agent", "Rejected path: release the reserved stock", "inventory_reservations RELEASED, audit_log", "-", "No PO, nothing committed"),
    (12, "Small amounts (up to $5,000)", "approve_budget commits at once (AUTO tier); POs created in the same run", "budget_ledger COMMIT, purchase_orders, audit_log", "create_purchase_order", "No human step"),
]

GOV = [
    ("Prompt", "The root instruction fixes the order of the steps; the specialists only return results to the root.", "campaign_provisioner/agent.py, sub_agents/"),
    ("Policy (data)", "Approval tiers and approver emails are BigQuery tables, changed without a release.", "approval_policy, approvers"),
    ("Platform", "The Application Integration approval pauses the flow until a person decides.", "request_approval trigger"),
    ("Code guard", "A purchase order is blocked unless an approved budget covers it (amount recomputed from the vendor price list).", "workflow/guard.py, v_request_headroom"),
    ("Code", "approve_budget rejects overspend; only the root can approve; the first approval decision wins (atomic update).", "tools/budget_tools.py, approval_flow.py"),
    ("Audit", "Every governed action is written to the BigQuery audit log.", "audit_log"),
    ("Output", "No commentary between steps: some chat front ends treat the first text as the end of the answer.", "Root and specialist instructions"),
]

HEAD_FILL = PatternFill("solid", fgColor="1F3A8A")
BAND = PatternFill("solid", fgColor="EEF3FF")
THIN = Side(style="thin", color="C9D3EA")


def sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for c in range(1, len(headers) + 1):
        h = ws.cell(row=1, column=c)
        h.font, h.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
        h.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if r % 2 == 1:
                cell.fill = BAND
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    return ws


SUMMARY = [
    ("Root Agent", "Campaign Provisioner Agent (Orchestrator)",
     "The root agent acts as a front-line coordinator by greeting users, collecting the campaign and the items, and delegating to the "
     "inventory, procurement and budget agents in a fixed order. It incorporates a critical human-in-the-loop approval gate: for spend "
     "above $5,000 it starts an Application Integration approval, tells the user that an approval email was sent, and holds all "
     "procurement until the approver clicks Approve. Once approved it commits the budget and has the purchase orders created; if "
     "rejected it releases the reserved stock.",
     "BigQuery, Application Integration", 150),
    ("Sub Agent", "Inventory Agent",
     "The inventory agent matches each requested item to the BigQuery catalog, checks free stock and reserves the available units for "
     "all items in one call. It records the shortfall per item that has to be procured, and releases the reserved stock when a "
     "request is rejected.",
     "BigQuery", 95),
    ("Sub Agent", "Procurement Agent",
     "The procurement agent quotes every shortfall line from the cheapest active vendor, weighing price and lead time, and reports the "
     "grand total. After the budget is approved it creates the purchase orders through Application Integration, behind a code guard "
     "that blocks any order not covered by an approved budget.",
     "BigQuery, Application Integration", 110),
    ("Sub Agent", "Budget Agent",
     "The budget agent checks the campaign's remaining budget and finds the approval tier (auto-approved up to $5,000, Marketing "
     "Director up to $50,000, VP Marketing and Finance Controller above) and the approver role in one call, then reports back to the "
     "root agent. It never approves spend or places orders.",
     "BigQuery", 95),
]


def summary_sheet(wb):
    """The one-table summary in the same layout as the container agent reference: no header row, application and owner
    merged over all rows, then agent type, agent name, description, data sources."""
    ws = wb.create_sheet("Agent Summary", 0)
    grey = Side(style="thin", color="BFBFBF")
    box = Border(top=grey, bottom=grey, left=grey, right=grey)
    for i, (kind, name, desc, src, height) in enumerate(SUMMARY, start=1):
        for col, val in enumerate(("Campaign Provisioner Agent (Google Next 2027)", OWNER, kind, name, desc, src), start=1):
            if i == 1 or col > 2:
                ws.cell(row=i, column=col, value=val)
        ws.row_dimensions[i].height = height
    n = len(SUMMARY)
    ws.merge_cells(start_row=1, start_column=1, end_row=n, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=n, end_column=2)
    for r in range(1, n + 1):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = box
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center" if c < 5 else "top",
                                       horizontal="center" if c in (3,) else "left")
    for c, w in enumerate([24, 14, 10, 20, 78, 22], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def main():
    wb = Workbook()
    wb.remove(wb.active)
    summary_sheet(wb)
    sheet(wb, "Agents", ["Application", "Owner", "Agent Type", "Agent Name", "Description", "Data sources and integrations"],
          [(APP, OWNER, t, n, d, s) for t, n, d, s in AGENTS], [34, 18, 13, 26, 90, 46])
    sheet(wb, "Tools", ["Agent", "Tool", "Kind", "What it does", "BigQuery reads", "BigQuery writes", "Application Integration", "Guardrail"],
          TOOLS, [18, 30, 26, 60, 38, 42, 34, 44])
    sheet(wb, "BigQuery", ["Table / view", "Type", "Purpose", "Read by", "Written by"], BQ, [26, 8, 62, 34, 40])
    sheet(wb, "Application Integration", ["Trigger", "Type", "Purpose", "Inputs", "Outputs", "Tasks", "Called by", "When"],
          AI, [24, 18, 34, 44, 34, 70, 44, 26])
    sheet(wb, "Flow (HITL example)", ["Step", "Who", "Action", "BigQuery", "Application Integration", "Result"],
          FLOW, [6, 26, 46, 52, 40, 52])
    sheet(wb, "Governance", ["Layer", "Rule", "Where"], GOV, [16, 90, 46])
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
